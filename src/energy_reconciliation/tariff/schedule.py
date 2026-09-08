"""Reading the band schedule out of the tariff workbook, and a synthetic stand-in.

The workbook is opened **read-only** and never modified. Only ``Sheet1`` carries data
(ANL-001 §2); ``Sheet2`` and ``Sheet3`` are empty.

**One difference from the consumption data worth knowing.** A CSV reading keeps its
source *text* exactly as written, because the source wrote text. Excel does not: a
datetime cell holds a number, and openpyxl turns it into a ``datetime``. So
``schedule_label_text`` here is **our rendering** of that value, not a preserved source
string. There is no source string to preserve.

**A regular schedule does not establish a timezone.** 17,520 labels at exactly 1800
seconds, with both 2013 clock-change hours present once each, shows the schedule is a
fixed nominal 48-slot grid with no daylight-saving representation. It says nothing
about the convention used by the *consumption* timestamps, and a unique schedule key
proves only that a join cannot multiply rows -- an arithmetic property, not semantic
time alignment.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Final

DEFAULT_WORKBOOK: Final[Path] = Path("data/raw/Tariffs.xlsx")
DATA_SHEET: Final[str] = "Sheet1"
EXPECTED_HEADERS: Final[tuple[str, str]] = ("TariffDateTime", "Tariff")
EXPECTED_STEP_SECONDS: Final[int] = 1800

WORKBOOK_SOURCE: Final[str] = "Tariffs.xlsx"
DEMO_SOURCE: Final[str] = "synthetic-demo"

#: A synthetic one-day schedule for the committed demo archive, so the scenario can be
#: built and checked without the real workbook. The data is INVENTED. Two slots differ
#: from Normal so the demo exercises more than one price.
#: Naive on purpose: schedule labels carry no timezone, and attaching one here would
#: invent the very convention the project records as unresolved.
DEMO_DAY: Final[datetime] = datetime(2013, 1, 1)  # noqa: DTZ001
DEMO_BANDS: Final[dict[str, str]] = {"00:30": "Low", "01:00": "High"}


class ScheduleError(ValueError):
    """The workbook does not have the shape the schedule model requires."""


@dataclass(frozen=True, slots=True)
class ScheduleRow:
    label: datetime
    band_label: str


@dataclass(frozen=True, slots=True)
class Schedule:
    """A validated band schedule plus the identity of what it was read from."""

    rows: tuple[ScheduleRow, ...]
    source: str
    source_sha256: str

    @property
    def first_label(self) -> datetime:
        return self.rows[0].label

    @property
    def last_label(self) -> datetime:
        return self.rows[-1].label

    @property
    def years(self) -> tuple[int, ...]:
        return tuple(sorted({r.label.year for r in self.rows}))

    @property
    def content_digest(self) -> str:
        """A digest of the labels and bands themselves.

        ``source_sha256`` identifies the *file*; this identifies the *schedule*. The
        scenario fingerprint uses both, so a schedule whose content differs always
        forces a rebuild even if something upstream reported the wrong file digest.
        """
        payload = "|".join(f"{r.label.isoformat()}={r.band_label}" for r in self.rows)
        return hashlib.sha256(payload.encode()).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1 << 20):
            digest.update(chunk)
    return digest.hexdigest()


def _validate(rows: list[ScheduleRow], source: str) -> tuple[ScheduleRow, ...]:
    """Refuse anything that would make a join unsafe or a band undefined.

    A **duplicated schedule label is rejected here**, before a single row reaches the
    warehouse. That is the whole defence against join multiplication: if one label
    could carry two bands, every consumption reading at that label would be counted
    twice, and the total would be silently too large.
    """
    if not rows:
        raise ScheduleError(f"{source}: no schedule rows found")
    rows.sort(key=lambda r: r.label)
    seen: dict[datetime, str] = {}
    for row in rows:
        if row.label in seen:
            raise ScheduleError(
                f"{source}: duplicated schedule label {row.label.isoformat()} "
                f"(bands {seen[row.label]!r} and {row.band_label!r}). A repeated label "
                "would multiply every consumption reading joined to it."
            )
        seen[row.label] = row.band_label
    off_grid = [
        r for r in rows if r.label.minute % 30 or r.label.second or r.label.microsecond
    ]
    if off_grid:
        raise ScheduleError(
            f"{source}: {len(off_grid)} schedule label(s) are not on the half-hour "
            f"grid, first {off_grid[0].label.isoformat()}"
        )
    blank = [r for r in rows if not r.band_label]
    if blank:
        raise ScheduleError(
            f"{source}: {len(blank)} schedule row(s) have no band label"
        )
    return tuple(rows)


def read_workbook(path: Path = DEFAULT_WORKBOOK) -> Schedule:
    """Read ``Sheet1`` of the tariff workbook. Read-only; the file is not modified."""
    from openpyxl import load_workbook

    digest = file_sha256(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if DATA_SHEET not in book.sheetnames:
            raise ScheduleError(f"{path}: no sheet named {DATA_SHEET}")
        sheet = book[DATA_SHEET]
        stream = sheet.iter_rows(values_only=True)
        header = next(stream, None)
        if header is None or tuple(str(c) for c in header[:2]) != EXPECTED_HEADERS:
            raise ScheduleError(
                f"{path}: expected headers {EXPECTED_HEADERS}, found {header!r}"
            )
        rows: list[ScheduleRow] = []
        for number, cells in enumerate(stream, start=2):
            label, band = (cells + (None, None))[:2]
            if label is None and band is None:
                continue
            if not isinstance(label, datetime):
                raise ScheduleError(
                    f"{path}!{DATA_SHEET}A{number}: expected a datetime, found {label!r}"
                )
            rows.append(ScheduleRow(label, str(band).strip()))
    finally:
        book.close()
    return Schedule(_validate(rows, str(path)), WORKBOOK_SOURCE, digest)


def demo_schedule() -> Schedule:
    """A synthetic 48-slot schedule for 2013-01-01. INVENTED, clearly labelled."""
    rows = []
    for slot in range(48):
        label = DEMO_DAY + timedelta(minutes=30 * slot)
        rows.append(ScheduleRow(label, DEMO_BANDS.get(f"{label:%H:%M}", "Normal")))
    payload = "|".join(f"{r.label.isoformat()}={r.band_label}" for r in rows)
    return Schedule(
        _validate(rows, DEMO_SOURCE),
        DEMO_SOURCE,
        hashlib.sha256(payload.encode()).hexdigest(),
    )


def step_report(schedule: Schedule) -> dict[str, int]:
    """Measured shape of the schedule: how many steps, and how many are half-hourly."""
    steps = [
        int((b.label - a.label).total_seconds())
        for a, b in zip(schedule.rows, schedule.rows[1:], strict=False)
    ]
    return {
        "rows": len(schedule.rows),
        "steps": len(steps),
        "steps_of_1800_seconds": sum(1 for s in steps if s == EXPECTED_STEP_SECONDS),
        "longest_step_seconds": max(steps) if steps else 0,
    }


def loaded_at() -> datetime:
    return datetime.now(UTC).replace(tzinfo=None)
