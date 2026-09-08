"""ANL-003 step 3: the tariff dimensions built by dbt, and what a failure leaves behind.

Three things are checked here, and the third is the one that is easy to get wrong.

**Types survive.** A price is money-adjacent, so ``DECIMAL`` is not a preference. The
Arrow schema is checked, but the Arrow schema is only a *claim*: what matters is what
DuckDB persisted, so ``information_schema.columns`` is read after a real build and the
values are queried back.

**Values are exactly the published ones.** The expectations below are written out by
hand from the publisher's figures rather than imported from ``prices.py``. Importing the
catalogue would make the test agree with the code by construction, which proves nothing.

**A refusal leaves the previous answer intact.** ``schedule._validate`` rejects a
duplicated label. The tests state precisely what that buys: on a fresh target no schedule
dimension is created at all, and where one already exists its rows and schema are
untouched. It does **not** buy whole-build atomicity, and the test that would claim so is
deliberately absent -- the independent models build regardless, and one of these tests
asserts exactly that so the limit is recorded rather than implied.
"""

from __future__ import annotations

import hashlib
import zipfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import duckdb
import pyarrow as pa
import pytest
from conftest import HEADER, MEMBER, row

from energy_reconciliation import dbt_run
from energy_reconciliation.ingest.loader import load_member
from energy_reconciliation.tariff import dimensions as dim
from energy_reconciliation.tariff.schedule import (
    ScheduleError,
    demo_schedule,
    loaded_at,
)

pytest.importorskip("dbt.cli.main", reason="dbt is not installed")

BUILD = dbt_run.BUILD_SCHEMA

#: The publisher's figures, written here by hand. See prices.CITATION_REPORT.
#: pence -> pounds is a division by 100 that must be exact at these scales.
EXPECTED_PRICES = {
    ("ToU", "High"): (
        Decimal("67.2000"),
        Decimal("0.672000"),
        "2013-01-01",
        "2014-01-01",
    ),
    ("ToU", "Normal"): (
        Decimal("11.7600"),
        Decimal("0.117600"),
        "2013-01-01",
        "2014-01-01",
    ),
    ("ToU", "Low"): (
        Decimal("3.9900"),
        Decimal("0.039900"),
        "2013-01-01",
        "2014-01-01",
    ),
    # The flat rate's period is UNKNOWN, so both bounds are NULL and it prices nothing.
    ("Std", "flat"): (Decimal("14.2280"), Decimal("0.142280"), None, None),
}


def _warehouse(tmp_path: Path, name: str) -> Path:
    """A minimal loaded warehouse: enough for the source contract, no dataset needed."""
    archive = tmp_path / f"{name}.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr(
            MEMBER,
            HEADER + row("MAC000001", "ToU", "2013-01-01 00:30:00.0000000", " 1 "),
        )
    database = tmp_path / f"{name}.duckdb"
    assert load_member(archive, MEMBER, database).complete
    return database


def _duplicate_label_workbook(path: Path) -> Path:
    """A workbook whose schedule repeats one label with two different bands."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(["TariffDateTime", "Tariff"])
    sheet.append([datetime(2013, 1, 1, 0, 0), "Normal"])  # noqa: DTZ001 - naive on purpose
    sheet.append([datetime(2013, 1, 1, 0, 30), "High"])  # noqa: DTZ001
    sheet.append([datetime(2013, 1, 1, 0, 30), "Low"])  # noqa: DTZ001 - the duplicate
    book.save(path)
    return path


def _build(database: Path, target: Path, *extra: str) -> int:
    return dbt_run.main(
        ["--database", str(database), *extra, "build", "--target-path", str(target)]
    )


def _schedule_snapshot(database: Path) -> tuple[str, list]:
    """Row digest and column layout of the schedule dimension, or a marker if absent."""
    con = duckdb.connect(str(database), read_only=True)
    try:
        columns = con.execute(
            "SELECT column_name, data_type, ordinal_position FROM information_schema.columns "
            f"WHERE table_schema = '{BUILD}' AND table_name = 'dim_tariff_band_schedule' "
            "ORDER BY ordinal_position"
        ).fetchall()
        if not columns:
            return "absent", []
        rows = con.execute(
            f"SELECT * FROM {BUILD}.dim_tariff_band_schedule ORDER BY schedule_label_naive"
        ).fetchall()
        digest = hashlib.sha256("|".join(str(r) for r in rows).encode()).hexdigest()
        return digest, columns
    finally:
        con.close()


# --------------------------------------------------------- the Arrow schema is declared
def test_the_arrow_schema_states_the_decimal_scales():
    """Stated, not inferred. A pandas object column would leave DuckDB to guess."""
    price = dim.PRICE_SCHEMA
    assert price.field("price_pence_per_kwh").type == pa.decimal128(9, 4)
    assert price.field("price_gbp_per_kwh").type == pa.decimal128(9, 6)
    assert price.field("effective_from").nullable
    assert price.field("effective_until").nullable


def test_the_arrow_price_table_carries_exact_decimals_not_floats():
    table = dim.price_table().to_pylist()
    by_key = {(r["tariff_group"], r["band_label"]): r for r in table}
    assert set(by_key) == set(EXPECTED_PRICES)
    for key, (pence, gbp, _, _) in EXPECTED_PRICES.items():
        assert by_key[key]["price_pence_per_kwh"] == pence
        assert by_key[key]["price_gbp_per_kwh"] == gbp
        assert isinstance(by_key[key]["price_gbp_per_kwh"], Decimal)


def test_the_schedule_table_records_which_source_it_came_from():
    """An invented schedule must be unmistakable in the data itself, on every row."""
    table = dim.schedule_table(demo_schedule(), loaded_at())
    sources = set(table.column("schedule_source").to_pylist())
    assert sources == {"synthetic-demo"}
    assert table.num_rows == 48


def test_an_unknown_schedule_variant_is_refused_never_defaulted():
    """Silently substituting invented data for the publisher's is the worst failure."""
    with pytest.raises(dim.ScheduleVariantError):
        dim.resolve_schedule("demoo")
    with pytest.raises(dim.ScheduleVariantError):
        dim.resolve_schedule("")


# ------------------------------------------------------------------ database preflight
def test_an_unset_database_is_refused(capsys):
    assert dbt_run.main(["--database", ""]) == 2
    assert "no database given" in capsys.readouterr().err


def test_a_nonexistent_database_is_refused_and_stays_nonexistent(tmp_path, capsys):
    """The whole point: dbt would have created it. Nothing may appear on disk."""
    target = tmp_path / "not-there.duckdb"
    assert dbt_run.main(["--database", str(target)]) == 2
    assert "does not exist" in capsys.readouterr().err
    assert not target.exists(), "the preflight must not create the database it refused"


def test_an_empty_file_is_refused_and_left_alone(tmp_path):
    target = tmp_path / "empty.duckdb"
    target.touch()
    assert dbt_run.main(["--database", str(target)]) == 2
    assert target.stat().st_size == 0


def test_a_database_without_readings_is_refused(tmp_path, capsys):
    target = tmp_path / "blank.duckdb"
    duckdb.connect(str(target)).close()
    assert dbt_run.main(["--database", str(target)]) == 2
    assert "no main.readings" in capsys.readouterr().err


def test_the_preflight_accepts_a_loaded_warehouse(tmp_path):
    database = _warehouse(tmp_path, "ok")
    assert dbt_run.validated_database(str(database)) == database


# ------------------------------------------------- what DuckDB actually persisted
@pytest.fixture(scope="module")
def built(tmp_path_factory) -> Path:
    """One real dbt build with the demo schedule, so the dataset is never needed."""
    tmp_path = tmp_path_factory.mktemp("dbt_dimensions")
    database = _warehouse(tmp_path, "dims")
    assert _build(database, tmp_path / "target", "--schedule", "demo") == 0
    return database


def test_the_persisted_price_columns_are_decimal_with_the_declared_scales(built):
    """information_schema, not the Arrow schema: this is what the database really has."""
    con = duckdb.connect(str(built), read_only=True)
    try:
        types = dict(
            con.execute(
                "SELECT column_name, data_type FROM information_schema.columns "
                f"WHERE table_schema = '{BUILD}' AND table_name = 'dim_tariff_price'"
            ).fetchall()
        )
    finally:
        con.close()
    assert types["price_pence_per_kwh"] == "DECIMAL(9,4)"
    assert types["price_gbp_per_kwh"] == "DECIMAL(9,6)"
    assert types["effective_from"] == "DATE"
    assert types["effective_until"] == "DATE"


def test_the_persisted_prices_are_the_published_figures(built):
    con = duckdb.connect(str(built), read_only=True)
    try:
        rows = con.execute(
            "SELECT tariff_group, band_label, price_pence_per_kwh, price_gbp_per_kwh, "
            "CAST(effective_from AS VARCHAR), CAST(effective_until AS VARCHAR), "
            "currency, price_unit "
            f"FROM {BUILD}.dim_tariff_price"
        ).fetchall()
    finally:
        con.close()
    assert len(rows) == len(EXPECTED_PRICES)
    for group, band, pence, gbp, valid_from, valid_until, currency, unit in rows:
        assert (pence, gbp, valid_from, valid_until) == EXPECTED_PRICES[(group, band)]
        assert (currency, unit) == ("GBP", "pence_per_kwh")
        # exact, in the database, with no float anywhere in the path
        assert gbp * 100 == pence


def test_the_flat_rate_keeps_null_validity_rather_than_an_invented_period(built):
    con = duckdb.connect(str(built), read_only=True)
    try:
        found = con.execute(
            f"SELECT effective_from IS NULL AND effective_until IS NULL FROM {BUILD}"
            ".dim_tariff_price WHERE band_label = 'flat'"
        ).fetchone()[0]
    finally:
        con.close()
    assert found is True


def test_the_demo_schedule_is_labelled_as_invented_in_every_row(built):
    con = duckdb.connect(str(built), read_only=True)
    try:
        sources = con.execute(
            f"SELECT DISTINCT schedule_source FROM {BUILD}.dim_tariff_band_schedule"
        ).fetchall()
        labels = con.execute(
            f"SELECT COUNT(*), COUNT(DISTINCT schedule_label_naive) FROM {BUILD}"
            ".dim_tariff_band_schedule"
        ).fetchone()
    finally:
        con.close()
    assert sources == [("synthetic-demo",)]
    assert labels == (48, 48)


def test_a_successful_build_records_the_dbt_versions_that_made_it(built):
    """Candidate-output identity, kept apart from the published scenario run."""
    con = duckdb.connect(str(built), read_only=True)
    try:
        record = con.execute(
            "SELECT dbt_core_version, dbt_duckdb_version, schedule_variant, target_schema "
            f"FROM {dbt_run.BUILD_RUN_TABLE}"
        ).fetchall()
        published = con.execute(
            "SELECT COUNT(*) FROM information_schema.tables "
            "WHERE table_schema = 'main' AND table_name = 'scenario_run'"
        ).fetchone()[0]
    finally:
        con.close()
    assert len(record) == 1
    versions = dbt_run.dbt_identity()
    assert record[0][:2] == (versions["dbt-core"], versions["dbt-duckdb"])
    assert record[0][2:] == ("demo", BUILD)
    assert published == 0, "no scenario_run is created or touched by a dbt build"


# ------------------------------------------------------- what a refusal leaves behind
def test_a_duplicate_label_fails_the_build_with_the_schedule_error(tmp_path):
    """The refusal that protects the join reaches the build, message intact."""
    database = _warehouse(tmp_path, "dup")
    workbook = _duplicate_label_workbook(tmp_path / "duplicate.xlsx")
    # the underlying refusal, at its own boundary
    with pytest.raises(ScheduleError, match="duplicated schedule label"):
        dim.resolve_schedule("workbook", workbook)
    # and through dbt
    assert (
        _build(
            database,
            tmp_path / "target",
            "--schedule",
            "workbook",
            "--workbook",
            str(workbook),
        )
        != 0
    )


def test_invalid_schedule_input_creates_no_schedule_dimension_on_a_fresh_target(
    tmp_path,
):
    database = _warehouse(tmp_path, "fresh")
    workbook = _duplicate_label_workbook(tmp_path / "duplicate.xlsx")
    assert (
        _build(
            database,
            tmp_path / "target",
            "--schedule",
            "workbook",
            "--workbook",
            str(workbook),
        )
        != 0
    )
    digest, columns = _schedule_snapshot(database)
    assert (digest, columns) == ("absent", []), (
        "a failed schedule build must leave no schedule dimension at all, "
        "not an empty or partial one"
    )
    con = duckdb.connect(str(database), read_only=True)
    try:
        statuses = con.execute(
            f"SELECT status FROM {dbt_run.BUILD_RUN_TABLE}"
        ).fetchall()
    finally:
        con.close()
    assert statuses == [(dbt_run.FAILED,)], (
        "the attempt is recorded, and recorded as failed -- never as a success"
    )


def test_a_failed_replacement_leaves_the_existing_dimension_untouched(tmp_path):
    """Rows *and* schema. A dimension that survived with a changed shape has not survived."""
    database = _warehouse(tmp_path, "existing")
    assert _build(database, tmp_path / "target", "--schedule", "demo") == 0
    before = _schedule_snapshot(database)
    assert before[0] != "absent"

    workbook = _duplicate_label_workbook(tmp_path / "duplicate.xlsx")
    assert (
        _build(
            database,
            tmp_path / "target",
            "--schedule",
            "workbook",
            "--workbook",
            str(workbook),
        )
        != 0
    )

    assert _schedule_snapshot(database) == before
    con = duckdb.connect(str(database), read_only=True)
    try:
        leftovers = con.execute(
            "SELECT table_name FROM information_schema.tables "
            "WHERE table_name LIKE '%dbt_tmp%' OR table_name LIKE '%backup%'"
        ).fetchall()
        attempts = con.execute(
            f"SELECT status FROM {dbt_run.BUILD_RUN_TABLE} ORDER BY started_at_utc"
        ).fetchall()
    finally:
        con.close()
    assert leftovers == [], "a failed build must not leave a temporary relation behind"
    assert attempts == [(dbt_run.SUCCEEDED,), (dbt_run.FAILED,)], (
        "both attempts are on record, and only the first as a success"
    )


def test_the_boundary_is_per_model_not_whole_build_atomicity(tmp_path):
    """Recorded as a limit, not implied.

    The price dimension has no dependency on the schedule one, so dbt builds it even
    though the schedule model failed. Anyone reading the guarantee above must know that
    "nothing was written" is true of the *failing model*, not of the run.
    """
    database = _warehouse(tmp_path, "partial")
    workbook = _duplicate_label_workbook(tmp_path / "duplicate.xlsx")
    assert (
        _build(
            database,
            tmp_path / "target",
            "--schedule",
            "workbook",
            "--workbook",
            str(workbook),
        )
        != 0
    )
    con = duckdb.connect(str(database), read_only=True)
    try:
        present = {
            name
            for (name,) in con.execute(
                "SELECT table_name FROM information_schema.tables "
                f"WHERE table_schema = '{BUILD}'"
            ).fetchall()
        }
    finally:
        con.close()
    assert "dim_tariff_band_schedule" not in present
    assert "dim_tariff_price" in present, (
        "independent models still build; this is why whole-build atomicity is not claimed"
    )
